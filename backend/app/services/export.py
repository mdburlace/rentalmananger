"""Accountant exports: Excel workbook, CSV and a ZIP pack with receipts."""
from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from ..ato import ATO_BY_KEY, ATO_LABELS, TRANSFER_KEYS
from ..config import get_settings
from ..fy import fy_label
from ..models import Portfolio, User
from ..money import apply_percent, from_cents
from .reports import fy_report, fy_transactions

MONEY = '"$"#,##0.00;[Red]-"$"#,##0.00'
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
SECTION_FILL = PatternFill("solid", fgColor="E8EEF5")
TOTAL_FILL = PatternFill("solid", fgColor="F3F4F6")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
THIN_TOP = Border(top=Side(style="thin", color="9CA3AF"))


def _autowidth(ws, min_w=10, max_w=60):
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                widths[c.column] = max(widths.get(c.column, 0), len(str(c.value)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_w, min(max_w, w + 2))


def _header(ws, row: int, values: list[str]):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = WHITE_BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)


def build_workbook(db: Session, user: User, fy: int, property_id: int | None = None) -> bytes:
    rep = fy_report(db, user, fy, property_id)
    portfolio = db.get(Portfolio, user.portfolio_id)
    props = rep["properties"]
    wb = Workbook()

    # ---- Summary: one column per property, rows follow the rental schedule
    ws = wb.active
    ws.title = "Rental schedule"
    ws["A1"] = f"Rental property schedule – FY {rep['fy_label']} (1 Jul {rep['start'].year} – 30 Jun {rep['end'].year})"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{portfolio.name} · generated {datetime.now():%d/%m/%Y %H:%M} · amounts are the claimable share after any private-use apportionment"
    ws["A2"].font = Font(italic=True, color="6B7280")
    n = len(props)
    multi = n > 1
    _header(ws, 4, ["", *[p["property"]["name"] for p in props], *(["All properties"] if multi else [])])
    r = 5

    def detail(label, values, fmt=None):
        nonlocal r
        ws.cell(row=r, column=1, value=label)
        for i, v in enumerate(values, start=2):
            c = ws.cell(row=r, column=i, value=v)
            if fmt:
                c.number_format = fmt
        r += 1

    def section(title):
        nonlocal r
        for col in range(1, n + 2 + (1 if multi else 0)):
            ws.cell(row=r, column=col).fill = SECTION_FILL
        ws.cell(row=r, column=1, value=title).font = BOLD
        r += 1

    section("Property details")
    detail("Address", [p["property"]["address"] for p in props])
    detail("Date purchased", [p["property"]["purchase_date"] for p in props], "dd/mm/yyyy")
    detail("Purchase price", [p["property"]["purchase_price"] for p in props], MONEY)
    detail("Date first earned rental income", [p["property"]["first_rented_date"] for p in props], "dd/mm/yyyy")
    detail("Weeks rented this year", [p["year"]["weeks_rented"] for p in props])
    detail("Weeks available for rent", [p["year"]["weeks_available"] for p in props])
    detail("Loan refinanced/renegotiated this year", ["Yes" if p["year"]["loan_refinanced"] else "No" for p in props])
    detail("Ownership", ["; ".join(f"{o['owner_name']} {o['share_percent']:g}%" for o in p["owners"]) for p in props])

    def money_rows(kind: str, title: str, total_label: str):
        nonlocal r
        section(title)
        first = r
        for a in ATO_LABELS:
            if a.kind != kind:
                continue
            ws.cell(row=r, column=1, value=a.label)
            for i, p in enumerate(props, start=2):
                line = next(l for l in p[{"income": "income", "expense": "expenses", "excluded": "excluded"}[kind]] if l["key"] == a.key)
                c = ws.cell(row=r, column=i, value=line["amount"] or 0)
                c.number_format = MONEY
            if multi:
                c = ws.cell(row=r, column=n + 2, value=f"=SUM(B{r}:{get_column_letter(n + 1)}{r})")
                c.number_format = MONEY
            r += 1
        ws.cell(row=r, column=1, value=total_label).font = BOLD
        for col in range(2, n + 2 + (1 if multi else 0)):
            L = get_column_letter(col)
            c = ws.cell(row=r, column=col, value=f"=SUM({L}{first}:{L}{r - 1})")
            c.number_format = MONEY
            c.font = BOLD
            c.fill = TOTAL_FILL
            c.border = THIN_TOP
        r += 1
        return r - 1

    inc_row = money_rows("income", "Income", "Total income")
    exp_row = money_rows("expense", "Expenses", "Total expenses")
    r += 1
    ws.cell(row=r, column=1, value="Net rent (income − expenses)").font = Font(bold=True, size=12)
    for col in range(2, n + 2 + (1 if multi else 0)):
        L = get_column_letter(col)
        c = ws.cell(row=r, column=col, value=f"={L}{inc_row}-{L}{exp_row}")
        c.number_format = MONEY
        c.font = Font(bold=True, size=12)
    r += 2
    money_rows("excluded", "Not claimed (capital / private) – for the CGT cost base", "Total not claimed")
    ws.freeze_panes = "B5"
    ws.column_dimensions["A"].width = 52
    for col in range(2, n + 3):
        ws.column_dimensions[get_column_letter(col)].width = 22

    # ---- Owner shares
    wo = wb.create_sheet("Owner shares")
    _header(wo, 1, ["Owner", "Property", "Share %", "Income", "Expenses", "Net rent"])
    rr = 2
    for p in props:
        for o in p["owners"]:
            vals = [o["owner_name"], p["property"]["name"], o["share_percent"] / 100, o["income"], o["expenses"], o["net"]]
            for i, v in enumerate(vals, start=1):
                c = wo.cell(row=rr, column=i, value=v)
                if i == 3:
                    c.number_format = "0.00%"
                elif i > 3:
                    c.number_format = MONEY
            rr += 1
    if rep["owners"]:
        rr += 1
        wo.cell(row=rr, column=1, value="Totals by owner").font = BOLD
        rr += 1
        for o in rep["owners"]:
            wo.cell(row=rr, column=1, value=o["owner_name"])
            for i, k in zip((4, 5, 6), ("income", "expenses", "net")):
                c = wo.cell(row=rr, column=i, value=o[k])
                c.number_format = MONEY
                c.font = BOLD
            rr += 1
    _autowidth(wo)

    # ---- Owner funding (contributions / drawings – not on the schedule)
    if rep["owner_funding"]:
        wf = wb.create_sheet("Owner funding")
        wf["A1"] = "Money owners put in or took out – not income or expenses, shown for information"
        wf["A1"].font = Font(italic=True, color="6B7280")
        _header(wf, 3, ["Property", "Owner", "Contributed", "Drawn", "Net put in"])
        rr = 4
        for p in props:
            for f in p["owner_funding"]:
                vals = [p["property"]["name"], f["owner_name"], f["contributed"], f["drawn"], f["net"]]
                for i, v in enumerate(vals, start=1):
                    c = wf.cell(row=rr, column=i, value=v)
                    if i > 2:
                        c.number_format = MONEY
                rr += 1
        if len(props) > 1:
            rr += 1
            wf.cell(row=rr, column=1, value="All properties").font = BOLD
            for f in rep["owner_funding"]:
                vals = ["", f["owner_name"], f["contributed"], f["drawn"], f["net"]]
                for i, v in enumerate(vals, start=1):
                    c = wf.cell(row=rr, column=i, value=v)
                    c.font = BOLD
                    if i > 2:
                        c.number_format = MONEY
                rr += 1
        _autowidth(wf)

    # ---- Transactions
    wt = wb.create_sheet("Transactions")
    cols = ["Date", "Property", "Schedule label", "Category", "Payee", "Description", "Reference", "Bill period",
            "Amount", "Claim %", "Claimable", "Receipts", "Needs review", "Notes", "Owner"]
    _header(wt, 1, cols)
    txns = fy_transactions(db, user, fy, property_id)
    for i, t in enumerate(txns, start=2):
        period = ""
        if t.period_start or t.period_end:
            period = f"{t.period_start:%d/%m/%Y}" if t.period_start else ""
            period += f" – {t.period_end:%d/%m/%Y}" if t.period_end else ""
        vals = [t.date, t.property.name, ATO_BY_KEY[t.category.ato_key].label, t.category.name, t.payee,
                t.description, t.reference, period, from_cents(t.amount_cents), float(t.deductible_percent) / 100,
                None if t.category.ato_key in TRANSFER_KEYS else from_cents(apply_percent(t.amount_cents, t.deductible_percent)),
                ", ".join(a.filename for a in t.attachments), "Yes" if t.needs_review else "", t.notes,
                t.owner.name if t.owner else None]
        for j, v in enumerate(vals, start=1):
            c = wt.cell(row=i, column=j, value=v)
            if j == 1:
                c.number_format = "dd/mm/yyyy"
            elif j in (9, 11):
                c.number_format = MONEY
            elif j == 10:
                c.number_format = "0%"
    last = len(txns) + 1
    if txns:
        wt.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
    wt.freeze_panes = "A2"
    _autowidth(wt)

    # ---- Loans
    wl = wb.create_sheet("Loans")
    _header(wl, 1, ["Property", "Loan", "Lender", "Account", "Rental use %", f"Balance at 30 Jun {rep['end'].year}",
                    "Balance date", "Interest rate %", "Interest claimed this FY"])
    rr = 2
    for p in props:
        interest = next(l for l in p["expenses"] if l["key"] == "interest")["amount"]
        for l in p["loans"]:
            vals = [p["property"]["name"], l["name"], l["lender"], l["account_hint"], l["rental_use_percent"] / 100,
                    l["balance"], l["balance_date"], l["interest_rate"], interest]
            for i, v in enumerate(vals, start=1):
                c = wl.cell(row=rr, column=i, value=v)
                c.number_format = {5: "0%", 6: MONEY, 7: "dd/mm/yyyy", 9: MONEY}.get(i, "General")
            rr += 1
    _autowidth(wl)

    # ---- Notes & checks
    wn = wb.create_sheet("Notes & checks")
    _header(wn, 1, ["Property", "Type", "Note"])
    rr = 2
    for p in props:
        if p["year"]["notes"]:
            for i, v in enumerate([p["property"]["name"], "Owner note", p["year"]["notes"]], start=1):
                wn.cell(row=rr, column=i, value=v)
            rr += 1
        for chk in p["checks"]:
            for i, v in enumerate([p["property"]["name"], chk["level"].title(), chk["message"]], start=1):
                wn.cell(row=rr, column=i, value=v)
            rr += 1
    wn.column_dimensions["A"].width = 24
    wn.column_dimensions["B"].width = 12
    wn.column_dimensions["C"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def transactions_csv(db: Session, user: User, fy: int, property_id: int | None = None) -> str:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["date", "property", "schedule_label", "category", "payee", "description", "reference",
                "period_start", "period_end", "amount", "claim_percent", "claimable", "receipts", "needs_review", "notes",
                "owner"])
    for t in fy_transactions(db, user, fy, property_id):
        w.writerow([t.date.isoformat(), t.property.name, ATO_BY_KEY[t.category.ato_key].label, t.category.name,
                    t.payee or "", t.description or "", t.reference or "",
                    t.period_start.isoformat() if t.period_start else "", t.period_end.isoformat() if t.period_end else "",
                    f"{t.amount_cents / 100:.2f}", f"{float(t.deductible_percent):g}",
                    f"{apply_percent(t.amount_cents, t.deductible_percent) / 100:.2f}",
                    "; ".join(a.filename for a in t.attachments), "yes" if t.needs_review else "", t.notes or "",
                    t.owner.name if t.owner else ""])
    return out.getvalue()


def _safe(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9._ -]+", "_", s).strip()[:80] or "file"


def accountant_pack(db: Session, user: User, fy: int, property_id: int | None = None) -> bytes:
    label = fy_label(fy)
    uploads = get_settings().uploads_dir
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"Rental schedule FY{label}.xlsx", build_workbook(db, user, fy, property_id))
        z.writestr(f"Transactions FY{label}.csv", transactions_csv(db, user, fy, property_id))
        count = 0
        for t in fy_transactions(db, user, fy, property_id):
            for a in t.attachments:
                path = uploads / a.stored_name
                if not path.exists():
                    continue
                name = f"{t.date:%Y-%m-%d}_{_safe(t.category.name)}_{t.amount_cents / 100:.2f}_{_safe(a.filename)}"
                z.write(path, f"Receipts/{_safe(t.property.name)}/{name}")
                count += 1
        z.writestr("README.txt", (
            f"Rental property records for FY {label}\n\n"
            f"- 'Rental schedule FY{label}.xlsx' – totals by ATO rental schedule label, owner shares, every "
            "transaction, loan balances and open items.\n"
            f"- 'Transactions FY{label}.csv' – the same transactions for importing into other software.\n"
            "- Owner contributions/drawings (if any) are listed on the 'Owner funding' sheet. They are not income\n"
            "  or expenses and are not included in the rental schedule totals.\n"
            f"- 'Receipts/' – {count} supporting document(s), named date_category_amount_original-name.\n"
        ))
    return buf.getvalue()
