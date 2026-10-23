"""Import transactions from spreadsheets (xlsx) or CSV exports.

Designed around the household's existing workbook – one sheet per financial
year with Date / Item / Value columns – but also handles typical bank CSVs
(Date / Description / Amount or Debit / Credit).

The import is two-step: `preview()` parses the file and proposes a category,
date and include/skip decision for each row (flagging missing dates, loan
balances, duplicates); the user adjusts and the client posts the rows back to
be committed.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..ato import TRANSFER_KEYS
from ..fy import fy_bounds, fy_from_text, fy_label, fy_of
from ..models import Category, Transaction, User
from ..money import from_cents, to_cents

DATE_HEADERS = {"date", "transaction date", "date paid", "paid", "posted", "value date"}
DESC_HEADERS = {"item", "description", "details", "narrative", "memo", "payee", "transaction details", "particulars", "expense"}
AMOUNT_HEADERS = {"value", "amount", "total", "cost", "amount ($)", "amount aud"}
DEBIT_HEADERS = {"debit", "debit amount", "withdrawal", "withdrawals", "money out"}
CREDIT_HEADERS = {"credit", "credit amount", "deposit", "deposits", "money in"}

DATE_FORMATS = ["%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d %b %Y", "%d %B %Y", "%d.%m.%Y"]


@dataclass
class PreviewRow:
    row_id: str
    sheet: str
    row_number: int
    raw_date: str | None
    description: str
    amount: float | None
    date: date | None
    fy: int | None
    category_id: int | None
    action: str  # transaction | loan_balance | skip
    needs_review: bool = False
    issues: list[str] = field(default_factory=list)
    duplicate_of: int | None = None
    interest_rate: float | None = None


def _norm(v) -> str:
    return str(v).strip().lower() if v is not None else ""


def _parse_date(v) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _read_tables(content: bytes, filename: str) -> list[tuple[str, list[list]]]:
    if filename.lower().endswith((".csv", ".txt")):
        text = content.decode("utf-8-sig", errors="replace")
        return [(filename, [row for row in csv.reader(io.StringIO(text))])]
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append((ws.title, [list(r) for r in ws.iter_rows(values_only=True)]))
    return out


def _find_header(rows: list[list]) -> tuple[int, dict[str, int]] | None:
    for i, row in enumerate(rows[:15]):
        cols: dict[str, int] = {}
        for j, cell in enumerate(row):
            h = _norm(cell)
            if h in DATE_HEADERS and "date" not in cols:
                cols["date"] = j
            elif h in DESC_HEADERS and "desc" not in cols:
                cols["desc"] = j
            elif h in AMOUNT_HEADERS and "amount" not in cols:
                cols["amount"] = j
            elif h in DEBIT_HEADERS:
                cols["debit"] = j
            elif h in CREDIT_HEADERS:
                cols["credit"] = j
        if "desc" in cols and ("amount" in cols or "debit" in cols or "credit" in cols):
            return i, cols
    return None


def _match_category(text: str, categories: list[Category], prefer_income: bool | None) -> int | None:
    t = text.lower()
    best: tuple[int, int] | None = None  # (keyword length, category id)
    for c in categories:
        if c.archived:
            continue
        kws = [k.strip().lower() for k in (c.keywords or "").split(",") if k.strip()]
        kws.append(c.name.lower())
        for k in kws:
            if k and re.search(r"(?<![a-z])" + re.escape(k), t):
                if best is None or len(k) > best[0]:
                    best = (len(k), c.id)
    return best[1] if best else None


_BALANCE_RE = re.compile(r"\bbalance\b", re.I)
_RATE_RE = re.compile(r"(?:rate|@)\s*([0-9]{1,2}(?:\.[0-9]+)?)\s*%?", re.I)


def preview(db: Session, user: User, property_id: int, content: bytes, filename: str) -> dict:
    categories = list(db.scalars(select(Category).where(Category.portfolio_id == user.portfolio_id)).all())
    cat_by_id = {c.id: c for c in categories}
    rows: list[PreviewRow] = []
    sheets_summary = []

    for sheet, table in _read_tables(content, filename):
        found = _find_header(table)
        if not found:
            sheets_summary.append({"sheet": sheet, "rows": 0, "skipped": "No Date/Item/Value header found"})
            continue
        hdr, cols = found
        sheet_fy = fy_from_text(sheet)
        count = 0
        for i, raw in enumerate(table[hdr + 1:], start=hdr + 2):
            def cell(key):
                j = cols.get(key)
                return raw[j] if j is not None and j < len(raw) else None

            desc = str(cell("desc") or "").strip()
            amount_raw = cell("amount")
            prefer_income = None
            if amount_raw in (None, "") and ("debit" in cols or "credit" in cols):
                if cell("credit") not in (None, ""):
                    amount_raw, prefer_income = cell("credit"), True
                elif cell("debit") not in (None, ""):
                    amount_raw, prefer_income = cell("debit"), False
            raw_date = cell("date")
            if not desc and amount_raw in (None, ""):
                continue
            count += 1
            r = PreviewRow(row_id=f"{sheet}:{i}", sheet=sheet, row_number=i,
                           raw_date=str(raw_date) if raw_date not in (None, "") else None,
                           description=desc, amount=None, date=None, fy=None, category_id=None, action="transaction")
            try:
                cents = to_cents(amount_raw)
                if cents < 0:
                    cents = -cents
                    prefer_income = False if prefer_income is None else prefer_income
                r.amount = from_cents(cents)
            except ValueError:
                r.issues.append(f"Amount '{amount_raw}' couldn't be read")
                r.action = "skip"

            d = _parse_date(raw_date)
            if raw_date not in (None, "") and d is None:
                r.issues.append(f"Date '{raw_date}' couldn't be read")
            if d is None and sheet_fy:
                d = fy_bounds(sheet_fy)[1]
                r.needs_review = True
                r.issues.append(f"No date – set to 30 Jun {sheet_fy} (end of FY {fy_label(sheet_fy)}). Enter the paid date if you have it.")
            elif d is None:
                r.needs_review = True
                r.issues.append("No date – please enter one")
            elif sheet_fy and fy_of(d) != sheet_fy:
                r.issues.append(f"Dated in FY {fy_label(fy_of(d))} but listed on the FY {fy_label(sheet_fy)} sheet")
                r.needs_review = True
            r.date = d
            r.fy = fy_of(d) if d else None

            if _BALANCE_RE.search(desc):
                r.action = "loan_balance"
                m = _RATE_RE.search(desc) or re.search(r"([0-9]{1,2}\.[0-9]+)", desc)
                if m:
                    r.interest_rate = float(m.group(1))
                r.issues.append("Looks like a loan balance, not an expense – will be saved as a loan balance")
            else:
                r.category_id = _match_category(desc, categories, prefer_income)
                if r.category_id is None and r.action == "transaction":
                    r.issues.append("Choose a category")
            rows.append(r)
        sheets_summary.append({"sheet": sheet, "fy": sheet_fy, "rows": count})

    _flag_duplicates(db, user, property_id, rows, cat_by_id)
    return {
        "filename": filename,
        "sheets": sheets_summary,
        "rows": [_row_dict(r) for r in rows],
    }


def _row_dict(r: PreviewRow) -> dict:
    d = asdict(r)
    d["date"] = r.date.isoformat() if r.date else None
    d["suggested_category_id"] = r.category_id
    return d


def _flag_duplicates(db: Session, user: User, property_id: int, rows: list[PreviewRow], cats: dict[int, Category]):
    """Match each incoming row to at most one existing transaction (same property,
    category and FY, amount within $1) so re-importing a file doesn't double up,
    and warn when the same bill amount lands in two financial years."""
    existing = list(db.scalars(select(Transaction).where(Transaction.property_id == property_id,
                                                         Transaction.portfolio_id == user.portfolio_id)).unique().all())
    used: set[int] = set()
    candidates = [r for r in rows if r.action == "transaction" and r.amount is not None and r.category_id and r.date]

    def try_match(r: PreviewRow, exact_date: bool, tolerance: int) -> bool:
        cents = to_cents(r.amount)
        for t in existing:
            if t.id in used or t.category_id != r.category_id or fy_of(t.date) != r.fy:
                continue
            if exact_date and t.date != r.date:
                continue
            if abs(t.amount_cents - cents) <= tolerance:
                used.add(t.id)
                r.duplicate_of = t.id
                r.action = "skip"
                r.issues.insert(0, f"Already recorded ({t.date:%d/%m/%Y} ${from_cents(t.amount_cents):,.2f}) – skipped")
                return True
        return False

    pending = list(candidates)
    for exact_date, tol in ((True, 0), (False, 0), (False, 100)):
        pending = [r for r in pending if not try_match(r, exact_date, tol)]

    # Within-file: identical rows (same date, amount, category) – keep the first.
    seen: dict[tuple, PreviewRow] = {}
    for r in candidates:
        if r.action != "transaction":
            continue
        key = (r.date, r.amount, r.category_id)
        if key in seen:
            r.action = "skip"
            r.issues.insert(0, f"Same as row {seen[key].row_number} on '{seen[key].sheet}' – skipped")
        else:
            seen[key] = r

    # Same amount & category in a different FY, close together → probably one bill counted twice.
    live = [r for r in candidates if r.action == "transaction"]
    for i, a in enumerate(live):
        for b in live[i + 1:]:
            if (a.category_id == b.category_id and a.amount == b.amount and a.fy != b.fy
                    and cats[a.category_id].ato_key not in ("interest", "gross_rent", *TRANSFER_KEYS)
                    and abs((a.date - b.date).days) < 200):
                for x, y in ((a, b), (b, a)):
                    x.needs_review = True
                    x.issues.append(f"Same ${x.amount:,.2f} also on '{y.sheet}' (FY {fy_label(y.fy)}) – "
                                    "is this the same bill counted in two years?")


def commit(db: Session, user: User, property_id: int, filename: str, rows: list[dict], learn: bool) -> dict:
    from ..models import ImportBatch, Loan, LoanBalance

    batch = ImportBatch(portfolio_id=user.portfolio_id, filename=filename, created_by=user.id)
    db.add(batch)
    db.flush()
    created = balances = 0
    cats = {c.id: c for c in db.scalars(select(Category).where(Category.portfolio_id == user.portfolio_id)).all()}
    from ..models import Owner
    owner_ids = set(db.scalars(select(Owner.id).where(Owner.portfolio_id == user.portfolio_id)).all())
    for r in rows:
        action = r.get("action")
        if action == "skip":
            continue
        d = date.fromisoformat(r["date"]) if r.get("date") else None
        if d is None or r.get("amount") is None:
            raise ValueError(f"Row {r.get('row_id')}: date and amount are required")
        if action == "loan_balance":
            loan = db.scalar(select(Loan).where(Loan.property_id == property_id).order_by(Loan.id))
            if not loan:
                loan = Loan(property_id=property_id, name="Investment loan")
                db.add(loan)
                db.flush()
            existing = db.scalar(select(LoanBalance).where(LoanBalance.loan_id == loan.id, LoanBalance.as_of == d))
            rate = Decimal(str(r["interest_rate"])) if r.get("interest_rate") is not None else None
            if existing:
                existing.balance_cents, existing.interest_rate = to_cents(r["amount"]), rate
            else:
                db.add(LoanBalance(loan_id=loan.id, as_of=d, balance_cents=to_cents(r["amount"]), interest_rate=rate))
            balances += 1
            continue
        cid = r.get("category_id")
        if cid not in cats:
            raise ValueError(f"Row {r.get('row_id')}: choose a category")
        notes = "; ".join(i for i in r.get("issues", []) if "No date" in i or "two years" in i or "sheet" in i) or None
        is_transfer = cats[cid].ato_key in TRANSFER_KEYS
        if r.get("owner_id") and r["owner_id"] not in owner_ids:
            raise ValueError(f"Row {r.get('row_id')}: unknown owner")
        if is_transfer and not r.get("owner_id"):
            notes = "; ".join(x for x in [notes, "Choose which owner"] if x)
        db.add(Transaction(
            portfolio_id=user.portfolio_id, property_id=property_id, category_id=cid, date=d,
            amount_cents=to_cents(r["amount"]), description=(r.get("description") or "")[:500] or None,
            needs_review=bool(r.get("needs_review")) or (is_transfer and not r.get("owner_id")),
            owner_id=r.get("owner_id") if is_transfer else None, notes=notes, source="import",
            import_batch_id=batch.id, created_by=user.id,
        ))
        created += 1
        if learn and r.get("description") and cid != r.get("suggested_category_id"):
            _learn_keyword(cats[cid], r["description"], cats.values())
    batch.row_count = created
    db.commit()
    return {"batch_id": batch.id, "transactions": created, "loan_balances": balances}


def _learn_keyword(cat: Category, description: str, all_cats) -> None:
    """Remember a description → category mapping when no category already claims it."""
    text = re.sub(r"\s+", " ", description.strip().lower())
    if not text or len(text) > 60 or re.search(r"\d", text):
        return
    for c in all_cats:
        kws = [k.strip() for k in (c.keywords or "").split(",") if k.strip()]
        if text in kws or c.name.lower() == text:
            return
    kws = [k.strip() for k in (cat.keywords or "").split(",") if k.strip()]
    kws.append(text)
    cat.keywords = ", ".join(kws)
