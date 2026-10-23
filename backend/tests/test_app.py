import io
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from app.fy import fy_from_text, fy_label, fy_of
from app.money import to_cents

from .conftest import cat_id

FIXTURE = Path(__file__).parent / "fixtures" / "sample_expenses.xlsx"


def test_fy_helpers():
    from datetime import date
    assert fy_of(date(2025, 7, 1)) == 2026
    assert fy_of(date(2026, 6, 30)) == 2026
    assert fy_label(2026) == "2025-26"
    assert fy_from_text("23 - 24yr expences") == 2024
    assert fy_from_text("2024-2025 Year") == 2025
    assert fy_from_text("Template") is None
    assert to_cents("$1,234.56") == 123456
    assert to_cents(2211.5) == 221150


def test_setup_only_once(client):
    r = client.post("/api/auth/setup", json={"portfolio_name": "x", "name": "x", "email": "x@x.com", "password": "password123"})
    assert r.status_code == 409


def test_ownership_must_total_100(client, owners):
    a, _ = owners
    r = client.post("/api/properties", json={"name": "Bad", "ownerships": [{"owner_id": a["id"], "share_percent": 60}]})
    assert r.status_code == 422


def test_transactions_and_report(client, prop, cats):
    rent = cat_id(cats, "Rent received")
    interest = cat_id(cats, "Interest on loan")
    for m in range(7, 13):
        assert client.post("/api/transactions", json={"property_id": prop["id"], "category_id": rent,
                                                      "date": f"2025-{m:02d}-15", "amount": 2000}).status_code == 201
    # Loan is 80% rental use -> only 80% claimable
    r = client.post("/api/transactions", json={"property_id": prop["id"], "category_id": interest, "date": "2025-08-07",
                                               "amount": "1000.00", "deductible_percent": 80})
    assert r.json()["claimable_amount"] == 800.0
    rep = client.get("/api/reports/fy", params={"fy": 2026}).json()
    p = next(x for x in rep["properties"] if x["property"]["id"] == prop["id"])
    assert p["totals"] == {"income": 12000.0, "expenses": 800.0, "net": 11200.0, "excluded": 0.0}
    assert [o["net"] for o in p["owners"]] == [5600.0, 5600.0]
    assert any("interest is missing" in c["message"] for c in p["checks"])


def test_attachment_and_pack(client, prop, cats):
    t = client.get("/api/transactions", params={"fy": 2026, "kind": "expense"}).json()["items"][0]
    r = client.post(f"/api/transactions/{t['id']}/attachments",
                    files=[("files", ("receipt.pdf", b"%PDF-1.4 test", "application/pdf"))])
    assert r.status_code == 201, r.text
    bad = client.post(f"/api/transactions/{t['id']}/attachments",
                      files=[("files", ("x.exe", b"MZ", "application/x-msdownload"))])
    assert bad.status_code == 400
    z = zipfile.ZipFile(io.BytesIO(client.get("/api/reports/fy/pack.zip", params={"fy": 2026}).content))
    names = z.namelist()
    assert any(n.startswith("Receipts/Test St/2025-08-07_Interest on loan_1000.00") for n in names)
    wb = load_workbook(io.BytesIO(z.read("Rental schedule FY2025-26.xlsx")))
    assert wb.sheetnames == ["Rental schedule", "Owner shares", "Transactions", "Loans", "Notes & checks"]


def test_import_preview_and_commit(client, prop, cats):
    with open(FIXTURE, "rb") as f:
        r = client.post("/api/imports/preview", data={"property_id": prop["id"]},
                        files={"file": ("sample.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    prev = r.json()
    rows = prev["rows"]
    by_desc = {}
    for x in rows:
        by_desc.setdefault(x["description"].strip(), []).append(x)
    # categories matched by keyword
    assert by_desc["Rates"][0]["category_id"] == cat_id(cats, "Council rates")
    assert by_desc["Youi House Insurance"][0]["category_id"] == cat_id(cats, "Landlord insurance")
    assert by_desc["CPU Pest Control"][0]["category_id"] == cat_id(cats, "Pest control")
    # undated row defaults to FY end and is flagged
    undated = [x for x in by_desc["Water"] if x["raw_date"] is None and x["sheet"].startswith("23")][0]
    assert undated["date"] == "2024-06-30" and undated["needs_review"]
    # loan balance row recognised
    bal = by_desc["Interest rate 7.04 Balance on Home Loan"][0]
    assert bal["action"] == "loan_balance" and bal["interest_rate"] == 7.04
    # date on the wrong year's sheet + same bill in two years are both flagged
    w = [x for x in by_desc["Water"] if x["sheet"].startswith("24")]
    assert any("two years" in i for x in w for i in x["issues"])
    gutter = by_desc["Gutter clean"][0]
    assert gutter["category_id"] == cat_id(cats, "Cleaning")
    gutter["category_id"] = cat_id(cats, "Repairs & maintenance")

    res = client.post("/api/imports/commit", json={"property_id": prop["id"], "filename": "sample.xlsx", "rows": rows})
    assert res.status_code == 200, res.text
    out = res.json()
    assert out["transactions"] == 10 and out["loan_balances"] == 1

    # Re-importing the same file finds everything already recorded.
    with open(FIXTURE, "rb") as f:
        again = client.post("/api/imports/preview", data={"property_id": prop["id"]},
                            files={"file": ("sample.xlsx", f, "application/octet-stream")}).json()
    tx_rows = [x for x in again["rows"] if x["action"] != "loan_balance"]
    assert all(x["action"] == "skip" for x in tx_rows)

    loans = client.get("/api/loans", params={"property_id": prop["id"]}).json()
    assert loans[0]["balances"][0]["balance"] == 367631.25

    # learned mapping: "gutter clean" now belongs to repairs
    repairs = next(c for c in client.get("/api/categories").json() if c["name"] == "Repairs & maintenance")
    assert "gutter clean" in repairs["keywords"]

    # undo
    batch = client.get("/api/imports").json()[0]
    assert client.delete(f"/api/imports/{batch['id']}").json()["deleted"] == 10


def test_viewer_cannot_edit(client, prop, cats):
    client.post("/api/users", json={"email": "acct@example.com", "name": "Accountant", "role": "viewer",
                                    "password": "password123"})
    from fastapi.testclient import TestClient
    viewer = TestClient(client.app)
    assert viewer.post("/api/auth/login", json={"email": "acct@example.com", "password": "password123"}).status_code == 200
    assert viewer.get("/api/reports/fy", params={"fy": 2026}).status_code == 200
    r = viewer.post("/api/transactions", json={"property_id": prop["id"], "category_id": cat_id(cats, "Water"),
                                               "date": "2025-09-01", "amount": 10})
    assert r.status_code == 403
    assert TestClient(client.app).get("/api/properties").status_code == 401
